from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "user-guide" / "assets"
OUTPUT = ROOT / "docs" / "user-guide" / "מדריך משתמש ReTrace.docx"

FOREST = "0B5D4B"
TEAL = "128C86"
BLUE = "2563EB"
AMBER = "D97706"
PALE = "EAF6F3"
PALE_BLUE = "EEF5FF"
PALE_GRAY = "F5F7F8"
LIGHT_BORDER = "D9D9D9"
RED = "DC2626"
INK = "17202A"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LIGHT_BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)
        borders.append(tag)


def rtl_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    return paragraph


def rtl_run(run, font="Arial", size=None, bold=None, color=None):
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font)
    rtl = run._element.get_or_add_rPr().find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        run._element.get_or_add_rPr().append(rtl)
    rtl.set(qn("w:val"), "1")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_text(doc, text, *, bold_lead=None, color=None, space_after=7, keep=False):
    p = rtl_paragraph(doc.add_paragraph())
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.keep_with_next = keep
    if bold_lead and text.startswith(bold_lead):
        rtl_run(p.add_run(bold_lead), bold=True, color=color or FOREST)
        rtl_run(p.add_run(text[len(bold_lead):]))
    else:
        rtl_run(p.add_run(text), color=color)
    return p


def add_step(doc, number, title, text):
    p = rtl_paragraph(doc.add_paragraph())
    p.paragraph_format.left_indent = Inches(.18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    rtl_run(p.add_run(f"{number}. {title} "), bold=True, color=FOREST)
    rtl_run(p.add_run(text))
    return p


def add_bullet(doc, text, *, color=TEAL):
    p = rtl_paragraph(doc.add_paragraph())
    p.paragraph_format.left_indent = Inches(.22)
    p.paragraph_format.space_after = Pt(4)
    rtl_run(p.add_run("• "), bold=True, color=color)
    rtl_run(p.add_run(text))
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    rtl_paragraph(p)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    rtl_run(p.add_run(text), bold=True, color="000000")
    return p


def add_caption(doc, text):
    p = rtl_paragraph(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(9)
    rtl_run(p.add_run(text), size=9, color="555555")
    return p


def add_image(doc, filename, caption, width=6.7):
    p = rtl_paragraph(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    add_caption(doc, caption)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tbl_pr.append(bidi)
    header_row_properties = table.rows[0]._tr.get_or_add_trPr()
    header_row_properties.append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, FOREST)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = rtl_paragraph(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        rtl_run(p.add_run(header), bold=True, color="FFFFFF", size=10.5)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else PALE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = rtl_paragraph(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_after = Pt(0)
            rtl_run(p.add_run(str(value)), size=10)
            if widths:
                cell.width = Inches(widths[index])
    after = rtl_paragraph(doc.add_paragraph())
    after.paragraph_format.space_after = Pt(2)
    return table


def add_page_break(doc):
    doc.add_page_break()


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rtl_paragraph(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    rtl_run(paragraph.add_run("ReTrace  |  מדריך משתמש  |  "), size=9, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(fld_char1)
    paragraph.add_run()._r.append(instr)
    paragraph.add_run()._r.append(fld_char2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.65)
section.bottom_margin = Inches(.62)
section.left_margin = Inches(.72)
section.right_margin = Inches(.72)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
for style_name, size in (("Title", 30), ("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12)):
    style = doc.styles[style_name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.rPr.rFonts.set(qn("w:cs"), "Arial")

for sec in doc.sections:
    add_page_number(sec.footer.paragraphs[0])

# Cover
p = rtl_paragraph(doc.add_paragraph(style="Title"), WD_ALIGN_PARAGRAPH.CENTER)
p.paragraph_format.space_before = Pt(80)
rtl_run(p.add_run("מדריך משתמש ReTrace"), bold=True, size=30, color="000000")
p = rtl_paragraph(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
rtl_run(p.add_run("מערכת ניהול מעקב ותחקור תרגילים"), bold=True, size=18, color=FOREST)
p.paragraph_format.space_after = Pt(26)
p = rtl_paragraph(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
rtl_run(p.add_run("מדריך מלא לאפליקציה ולמערכת הצפייה"), size=14, color=TEAL)
p.paragraph_format.space_after = Pt(44)
for label, color in (("אפליקציה", TEAL), ("ניהול תרגילים", FOREST), ("מפה וציר זמן", BLUE), ("אירועים וייצוא", AMBER)):
    p = rtl_paragraph(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    rtl_run(p.add_run(f"●  {label}"), bold=True, size=13, color=color)
add_text(doc, "גרסה מעודכנת: אוקטובר 2026", color="666666", space_after=0).alignment = WD_ALIGN_PARAGRAPH.CENTER

add_page_break(doc)
heading(doc, "מטרת המדריך", 1)
add_text(doc, "המדריך מיועד למשתמש חדש שלא עבד בעבר עם ReTrace. הוא מסביר כיצד להיכנס, לבחור תפקיד, ליצור תרגיל או להצטרף לתרגיל, לעקוב אחר מיקום, לדווח אירוע, לצפות במסלולים, לנהל משתמשים ולייצא וידאו.")
add_text(doc, "המדריך מחולק לשני חלקים ברורים: האפליקציה בטלפון, המשמשת את המשתתפים בשטח; והמערכת לצפייה בתרגילים בדפדפן, המשמשת לצפייה, ניהול ותחקור.")
heading(doc, "איך להשתמש במדריך", 2)
add_bullet(doc, "בצילומי המסך, עיגול אדום מסמן את הכפתור שעליו צריך ללחוץ.", color=RED)
add_bullet(doc, "פקדים מסוימים מוצגים רק לבעלי הרשאה מתאימה או רק בתרגיל פעיל.")
add_bullet(doc, "המונח תפקיד בתרגיל שונה מסוג המשתמש. ההבדל מוסבר בפרק הבא.")
heading(doc, "תוכן העניינים", 2)
for text in [
    "1  סוגי משתמשים והרשאות", "2  חלק ראשון האפליקציה בטלפון", "3  חלק שני מערכת הצפייה בדפדפן",
    "4  תרחישים נפוצים", "5  פתרון תקלות", "6  כללי עבודה מומלצים ומילון מונחים"
]:
    add_bullet(doc, text)

add_page_break(doc)
heading(doc, "1 סוגי משתמשים והרשאות", 1)
add_text(doc, "לכל חשבון יש סוג משתמש קבוע במערכת. בנוסף, בכל הצטרפות לתרגיל המשתמש בוחר תפקיד מבצעי. סוג המשתמש קובע מה מותר לנהל; התפקיד המבצעי קובע כיצד המשתתף מוצג בתרגיל והאם הוא יכול לדווח אירועים.")
add_table(doc,
          ["סוג משתמש", "מה הוא יכול לעשות", "מה הוא לא יכול לעשות"],
          [
              ["ADMIN", "גישה מלאה לכל התרגילים; יצירה, שינוי שם, סגירה ומחיקה; ניהול משתמשים והרשאות; שיוך גישה לתרגילים.", "אין מגבלות מערכתיות רגילות."],
              ["MANAGER", "יצירה, שינוי שם וסגירה של תרגילים שאליהם שויך; צפייה בנתוני המשתתפים; ניהול גישה לתרגיל; דיווח אירועים.", "אינו יכול למחוק תרגיל לצמיתות ואינו יכול לנהל משתמשים או לשנות סיווג משתמש."],
              ["USER", "הצטרפות לתרגיל פעיל שאליו יש גישה; מעקב מיקום; צפייה בנתונים האישיים בלבד באתר.", "אינו יוצר או סוגר תרגילים ואינו צופה בנתונים של משתתפים אחרים."],
          ], widths=[1.15, 3.05, 2.65])
heading(doc, "תפקידים בתוך תרגיל", 2)
add_table(doc,
          ["סוג חשבון", "תפקידים שניתן לבחור", "ברירת מחדל"],
          [
              ["USER", "כיתת כוננות או חמל", "כיתת כוננות"],
              ["MANAGER", "רבשץ, כיתת כוננות, חמל או מנהל תרגיל", "רבשץ"],
              ["ADMIN", "רבשץ, כיתת כוננות, חמל או מנהל תרגיל", "מנהל תרגיל"],
          ], widths=[1.25, 3.7, 1.85])
add_text(doc, "דיווח אירועים: משתתף בתפקיד כיתת כוננות רואה את מסך המעקב ללא אזור דיווח אירוע. רבשץ, חמל ומנהל תרגיל יכולים לדווח אירועים. משתמש USER שבוחר חמל יכול לדווח אירועים במסגרת התרגיל.", bold_lead="דיווח אירועים:")
heading(doc, "כיצד קובעים סיווג משתמש", 2)
add_step(doc, 1, "התחבר כ ADMIN.", "ניהול המשתמשים מוצג בצד מערכת הצפייה.")
add_step(doc, 2, "הוסף משתמש חדש.", "הזן מייל, סיסמה זמנית ובחר USER, MANAGER או ADMIN.")
add_step(doc, 3, "שנה סיווג קיים.", "ברשימת המשתמשים בחר את הסיווג החדש. השינוי נשמר במערכת.")
add_text(doc, "חשוב: מומלץ להחזיק לפחות שני חשבונות ADMIN פעילים, כדי למנוע מצב שבו אין מנהל זמין למערכת.", bold_lead="חשוב:", color=RED)

add_page_break(doc)
heading(doc, "2 חלק ראשון האפליקציה בטלפון", 1)
heading(doc, "2.1 כניסה לאפליקציה", 2)
add_step(doc, 1, "פתח את ReTrace.", "יופיע מסך כניסה עם שדות כתובת מייל וסיסמה.")
add_step(doc, 2, "הזן את פרטי החשבון.", "יש להשתמש במייל ובסיסמה שהוגדרו על ידי מנהל המערכת.")
add_step(doc, 3, "לחץ על כניסה.", "לאחר אימות מוצלח יופיע מסך בחירת התרגיל.")
add_text(doc, "אם הסיסמה נשכחה, לחץ על שכחתי סיסמה, הזן כתובת מייל ולחץ על שלח קישור. הקישור יישלח רק אם הכתובת רשומה במערכת.", bold_lead="אם הסיסמה נשכחה,")
add_text(doc, "האפליקציה משתמשת באותו חשבון כמו מערכת הצפייה באתר.")

heading(doc, "2.2 פרטי המשתתף ובחירת תפקיד", 2)
add_text(doc, "בראש מסך ההגדרה מופיעים פרטי החשבון. בשדה שם הזן את השם שיוצג במפה ובתחקיר. לאחר מכן בחר תפקיד מתוך הרשימה המותרת לחשבון.")
add_bullet(doc, "השתמש בשם ברור ומלא, למשל נועה לוי או צוות סיור 1.")
add_bullet(doc, "אל תבחר תפקיד שאינו משקף את הפעילות בפועל.")
add_bullet(doc, "אם חסר תפקיד נדרש, פנה ל-ADMIN לבדיקת סוג החשבון.")

add_page_break(doc)
heading(doc, "2.3 יצירת תרגיל חדש", 2)
add_text(doc, "האפשרות זמינה לחשבונות ADMIN ו-MANAGER. משתמש USER אינו יכול ליצור תרגיל חדש.")
add_step(doc, 1, "הזן שם תרגיל.", "בחר שם קצר ומזהה, למשל תרגיל מוכנות יישובי 18.10.")
add_step(doc, 2, "לחץ על צור והתחל.", "המערכת יוצרת תרגיל פעיל, מוסיפה אותך כמשתתף ופותחת את מסך המעקב.")
add_image(doc, "mobile-setup-create.png", "מסך האפליקציה ליצירת תרגיל. הכפתור צור והתחל מסומן באדום.", width=3.55)

add_page_break(doc)
heading(doc, "2.4 הצטרפות לתרגיל פעיל", 2)
add_text(doc, "כל משתמש יכול להצטרף רק לתרגיל פעיל שאליו קיימת לו גישה. תרגיל שנסגר אינו מופיע ברשימה.")
add_step(doc, 1, "פתח את רשימת התרגילים הפעילים.", "אם הרשימה אינה מעודכנת, לחץ על כפתור הרענון.")
add_step(doc, 2, "בחר את התרגיל הנכון.", "בדוק את השם ואת שעת ההתחלה כדי להימנע מהצטרפות לתרגיל אחר.")
add_step(doc, 3, "לחץ על הצטרף והתחל מעקב.", "אשר הרשאת מיקום אם מערכת ההפעלה מבקשת זאת.")
add_image(doc, "mobile-setup-join.png", "מסך האפליקציה להצטרפות לתרגיל פעיל. הכפתור הצטרף והתחל מעקב מסומן באדום.", width=3.55)

add_page_break(doc)
heading(doc, "2.5 מסך המעקב", 2)
add_text(doc, "לאחר ההצטרפות האפליקציה מתחילה לאסוף מיקומים. באנדרואיד המעקב ממשיך גם כאשר עוברים לאפליקציה אחרת, כל עוד הרשאות המיקום וההתראות פעילות והמערכת אינה מגבילה את ReTrace ברקע.")
add_table(doc,
          ["רכיב", "משמעות"],
          [
              ["מקליט GPS", "הטלפון אוסף מיקומים עבור התרגיל."],
              ["זמן מעקב", "משך הזמן מאז תחילת המעקב הנוכחי."],
              ["GPS", "מציג את הדיוק המשוער של המיקום האחרון."],
              ["שרת", "מחובר ומסונכרן או Offline וינסה שוב."],
              ["נקודות שנשמרו", "מספר הנקודות שנשמרו עבור התרגיל הנוכחי בלבד."],
              ["ממתינות לסנכרון", "נקודות שנשמרו בטלפון ועדיין לא הועברו לשרת."],
          ], widths=[2.0, 4.7])
add_text(doc, "מצב Offline אינו גורם לאובדן מיידי של נתונים. הנקודות נשמרות בטלפון ונשלחות אוטומטית כאשר החיבור חוזר.", bold_lead="מצב Offline")

heading(doc, "2.6 עצירה והפעלה מחדש של המעקב", 2)
add_text(doc, "לחיצה על עצור מעקב מפסיקה את איסוף המיקומים. לאחר העצירה הכפתור משתנה להפעל מחדש. אין צורך לצאת מהתרגיל כדי לבצע הפסקה זמנית.")
add_image(doc, "mobile-tracking-stop.png", "מסך המעקב. הכפתור עצור מעקב מסומן באדום ונמצא לצד זמן המעקב.", width=3.55)

add_page_break(doc)
heading(doc, "2.7 דיווח אירוע מהאפליקציה", 2)
add_text(doc, "אזור דיווח האירוע מופיע עבור רבשץ, חמל ומנהל תרגיל. כל הפקדים נמצאים במסגרת ירוקה כדי להבהיר שהם חלק מתהליך אחד.")
add_step(doc, 1, "כתוב תיאור.", "תאר בקצרה מה קרה. מומלץ לכלול פעולה או השפעה ולא רק כותרת.")
add_step(doc, 2, "בדוק תאריך ושעה.", "ברירת המחדל היא הזמן הנוכחי. ניתן לבחור זמן אחר שאינו מוקדם מיצירת התרגיל ואינו מאוחר מהזמן הנוכחי.")
add_step(doc, 3, "בחר מיקום.", "לחץ על הוסף מיקום. אפשר לבחור מיקום עצמי או ללחוץ על נקודה במפה ולאשר.")
add_step(doc, 4, "לחץ על הוסף אירוע.", "שם המדווח, תפקידו, הזמן והמיקום נשמרים אוטומטית עם התיאור.")
add_image(doc, "mobile-event-report.png", "אזור דיווח האירוע באפליקציה. הכפתור הוסף אירוע מסומן באדום.", width=3.55)

add_page_break(doc)
heading(doc, "2.8 סיום תרגיל ויציאה", 2)
add_text(doc, "ADMIN או MANAGER יכולים לשנות את שם התרגיל או לסגור אותו. לפני סגירה מוצג חלון אישור. לאחר הסגירה התרגיל אינו מופיע ברשימת התרגילים הפעילים ולא ניתן להצטרף אליו.")
add_text(doc, "אם תרגיל נסגר בזמן שמשתתפים עוקבים, האפליקציה מנסה לסנכרן את הנקודות האחרונות, מפסיקה את המעקב ומציגה הודעה. אם התרגיל נמחק, המעקב נעצר ולא ניתן לשלוח אליו נתונים נוספים.")
add_text(doc, "בתפריט החשבון ניתן לבחור יציאה מהמכשיר הזה או יציאה מכל המכשירים. יציאה מכל המכשירים מנתקת את כל ההתקנות והדפדפנים לאחר שתוקף הגישה הקיימת מסתיים או בעת הרענון הבא.")
heading(doc, "הרשאות טלפון מומלצות", 2)
add_bullet(doc, "אפשר מיקום מדויק.")
add_bullet(doc, "אפשר מיקום ברקע באנדרואיד כאשר האפשרות מוצגת.")
add_bullet(doc, "אפשר התראות, כדי שמעקב הרקע יוצג כשירות פעיל.")
add_bullet(doc, "בטל הגבלת סוללה עבור ReTrace אם הטלפון מפסיק לאסוף נקודות ברקע.")

add_page_break(doc)
heading(doc, "3 חלק שני מערכת הצפייה בדפדפן", 1)
heading(doc, "3.1 כניסה לאתר", 2)
add_text(doc, "פתח את כתובת מערכת הצפייה בדפדפן Chrome או Edge. הזן את אותו מייל ואותה סיסמה המשמשים באפליקציה ולחץ על כניסה.")
add_image(doc, "web-login-action.png", "מסך הכניסה למערכת הצפייה. כפתור כניסה מסומן באדום.", width=6.5)
add_text(doc, "לאחר הכניסה, המידע המוצג תלוי בסוג החשבון ובגישה שניתנה לתרגיל. USER רואה רק את נתוניו; MANAGER רואה תרגילים שאליהם שויך; ADMIN רואה ומנהל את כלל המערכת.")

add_page_break(doc)
heading(doc, "3.2 בחירת תרגיל והצגת מסלולים", 2)
add_step(doc, 1, "בחר תרגיל מהרשימה העליונה.", "ליד שם התרגיל מופיע הסטטוס ACTIVE, ENDING או COMPLETED.")
add_step(doc, 2, "סמן משתתפים.", "סמן את תיבות הבחירה ליד כל משתתף שברצונך להציג. ניתן להציג כמה מסלולים יחד.")
add_step(doc, 3, "בדוק את המפה.", "לכל משתתף צבע משלו. אדום שמור לאירועים בלבד.")
add_text(doc, "ראש החץ נמצא במקום הנקודה האחרונה ומראה את כיוון התנועה. אם הצג שמות משתתפים מסומן, השם מוצג מעל החץ; אחרת השם מופיע בעת ריחוף.")
add_image(doc, "web-dashboard.png", "מערכת הצפייה עם שני מסלולים, נקודות, חיצי כיוון ושמות משתתפים.", width=6.9)

add_page_break(doc)
heading(doc, "3.3 מפה ציר זמן וניגון", 2)
add_table(doc,
          ["פקד", "פעולה"],
          [
              ["PLAY", "מתחיל או ממשיך ניגון מנקודת הזמן הנוכחית."],
              ["PAUSE", "עוצר זמנית את ההתקדמות באותה נקודת זמן."],
              ["STOP", "עוצר את הניגון ומחזיר את הזמן להתחלה. גם יוצא ממצב צפייה חיה."],
              ["קצב", "1, 5, 10 או 50 נקודות בשנייה, או ערך מותאם אישית בין 1 ל-100."],
              ["שובל", "מציג רק את 10, 30, 50 או 100 הנקודות האחרונות לכל משתתף."],
              ["הצג אירועים", "מוסיף אירועים לציר הזמן ומציג אותם במפה כשהזמן מגיע אליהם."],
              ["הצג שמות משתתפים", "מציג את השמות באופן קבוע מעל חצי הכיוון."],
          ], widths=[1.65, 5.05])
add_text(doc, "אפשר לגרור את סמן ציר הזמן או ללחוץ על נקודה בציר. בריחוף מוצגת השעה שאליה עומדים לעבור.")
add_image(doc, "web-dashboard-playback.png", "פקדי הניגון וציר הזמן בתחתית המפה מסומנים באדום.", width=6.9)

add_page_break(doc)
heading(doc, "3.4 סוגי מפה", 2)
add_bullet(doc, "מפה: רחובות, שמות יישובים ומקומות מרכזיים.")
add_bullet(doc, "לוויין: צילום אוויר ללא שכבת הרחובות המלאה.")
add_bullet(doc, "תצוגה משולבת: צילום לוויין עם שכבת שמות ורחובות מעליו.")
add_text(doc, "בחר את שכבת המפה מכפתור השכבות בפינת המפה. הבחירה נשמרת בדפדפן ומשמשת גם כברירת מחדל לייצוא וידאו.")
heading(doc, "3.5 אירועים על המפה", 2)
add_text(doc, "אירועים מוצגים תמיד באדום ובצורת כוכב או פיצוץ. בחמש השניות הראשונות לאחר הופעתם הם מהבהבים והתווית מוצגת קבוע. לאחר מכן הסמל נשאר במפה והתווית מופיעה בלחיצה.")
add_text(doc, "אירועים אינם נעלמים כאשר שובל מופעל. ביטול הצג אירועים מסיר אותם מהמפה ומציר הזמן.")
heading(doc, "3.6 צפייה חיה", 2)
add_text(doc, "צפייה חיה זמינה רק לתרגיל שעדיין פעיל. בלחיצה מוצג המיקום האחרון שנשמר לכל משתתף שנבחר. הגדרות שובל ואירועים ממשיכות לפעול. כדי לחזור לצפייה לאחור, לחץ על STOP.")

add_page_break(doc)
heading(doc, "3.7 הוספת אירוע מהאתר", 2)
add_text(doc, "הכפתור הוסף אירוע מופיע רק ל-ADMIN או MANAGER, רק בתרגיל פעיל ורק לאחר מעבר לצפייה חיה.")
add_step(doc, 1, "לחץ על צפייה חיה.", "לאחר המעבר יופיע הכפתור הוסף אירוע.")
add_step(doc, 2, "הזן תיאור ותאריך.", "התאריך חייב להיות בתוך טווח התרגיל ועד הזמן הנוכחי.")
add_step(doc, 3, "לחץ על המפה.", "בחר את הנקודה המדויקת שבה התרחש האירוע.")
add_step(doc, 4, "לחץ על שמור אירוע.", "האירוע יופיע בציר הזמן ובמפה בהתאם לזמנו.")
add_image(doc, "web-event-save.png", "חלון הוספת אירוע באתר. לאחר מילוי הנתונים ובחירת נקודה, לחץ על שמור אירוע.", width=6.9)

add_page_break(doc)
heading(doc, "3.8 ייצוא תרגיל לווידאו", 2)
add_text(doc, "ייצוא הווידאו נוצר בדפדפן במחשב, ללא עומס על שרת Render. בזמן ההכנה יש להשאיר את הלשונית פתוחה ופעילה.")
add_step(doc, 1, "בחר את המסלולים הרצויים.", "אפשר לבחור משתתפים שונים מאלה שמוצגים כעת במפה.")
add_step(doc, 2, "לחץ על ייצוא וידאו.", "הכפתור נמצא בסרגל העליון וזמין כאשר קיימות נקודות מיקום.")
add_image(doc, "web-dashboard-export.png", "כפתור ייצוא וידאו בסרגל העליון מסומן באדום.", width=6.9)
add_step(doc, 3, "הגדר את הסרטון.", "בחר סוג מפה, איכות, קצב, שובל, אירועים, שמות משתתפים וטווח זמן.")
add_step(doc, 4, "לחץ על צור וידאו.", "התקדמות היצירה תוצג בחלון. בסיום יורד קובץ WebM למחשב.")
add_image(doc, "web-export-create.png", "חלון הגדרות הייצוא. כפתור צור וידאו מסומן באדום.", width=6.9)
add_text(doc, "אם ההורדה אינה מתחילה, בדוק שהדפדפן מאפשר הורדות ושהלשונית נשארה פתוחה. מומלץ להשתמש ב-Chrome או Edge מעודכן.", bold_lead="אם ההורדה אינה מתחילה,")

add_page_break(doc)
heading(doc, "3.9 ניהול תרגילים", 2)
add_table(doc,
          ["פעולה", "מי רשאי", "תוצאה"],
          [
              ["תרגיל חדש", "ADMIN או MANAGER", "יוצר תרגיל פעיל ומשייך את היוצר אליו."],
              ["שנה שם", "ADMIN או MANAGER משויך", "מעדכן את שם התרגיל ללא שינוי הנתונים."],
              ["סגור תרגיל", "ADMIN או MANAGER משויך", "מונע הצטרפות חדשה ומסיים את המעקב הפעיל."],
              ["מחק תרגיל", "ADMIN בלבד", "מוחק לצמיתות את התרגיל, המשתתפים, הנקודות והאירועים."],
          ], widths=[1.35, 2.15, 3.2])
add_text(doc, "מחיקת תרגיל היא בלתי הפיכה. לפני המחיקה מוצג חלון אזהרה ויש לסמן תיבת אישור שבה מופיע שם התרגיל. בדוק את השם לפני האישור.", bold_lead="מחיקת תרגיל היא בלתי הפיכה.", color=RED)
heading(doc, "3.10 גישה לתרגיל", 2)
add_text(doc, "ADMIN או MANAGER משויך יכולים להוסיף חשבון קיים לתרגיל לפי כתובת מייל. הרשאת הגישה לתרגיל אינה משנה את סוג המשתמש הכללי.")
add_bullet(doc, "MANAGER שקיבל גישה יכול לנהל את אותו תרגיל בלבד.")
add_bullet(doc, "USER שקיבל גישה רואה רק את המסלולים והאירועים השייכים לחשבונו.")
add_bullet(doc, "בתרגילים ישנים אפשר לשייך משתתף היסטורי לחשבון לאחר אימות זהותו.")
add_bullet(doc, "הסר גישה חוסם גישה עתידית אך אינו מוחק את ההיסטוריה.")

add_page_break(doc)
heading(doc, "3.11 ניהול משתמשים", 2)
add_text(doc, "אזור ניהול המשתמשים מוצג ל-ADMIN בלבד.")
add_step(doc, 1, "הוספת משתמש.", "הזן כתובת מייל, סיסמה זמנית ובחר סוג משתמש.")
add_step(doc, 2, "שינוי הרשאה.", "בחר USER, MANAGER או ADMIN מהרשימה ליד המשתמש.")
add_step(doc, 3, "מחיקת משתמש.", "לחץ על מחק ואשר. לפני המחיקה בדוק שאין צורך לשמר גישה פעילה.")
add_text(doc, "MANAGER אינו יכול לראות את אזור ניהול המשתמשים ואינו יכול לשנות סיווגים. USER מקבל ממשק צפייה מוגבל בלבד.")

add_page_break(doc)
heading(doc, "4 תרחישים נפוצים", 1)
heading(doc, "תרחיש א משתתף כיתת כוננות", 2)
add_step(doc, 1, "התחבר באפליקציה כ USER.", "ברירת המחדל לתפקיד היא כיתת כוננות.")
add_step(doc, 2, "בחר תרגיל פעיל והצטרף.", "אשר הרשאות מיקום והמתן למצב מקליט GPS.")
add_step(doc, 3, "השאר את המעקב פעיל.", "ניתן לעבור לאפליקציות אחרות. בדוק מדי פעם שאין נקודות רבות שממתינות לסנכרון.")
add_step(doc, 4, "בסיום עצור מעקב.", "ודא שהנקודות הסתנכרנו לפני יציאה או כיבוי המכשיר.")
heading(doc, "תרחיש ב מנהל פותח ומנהל תרגיל", 2)
add_step(doc, 1, "צור תרגיל חדש.", "בחר שם ברור והתחל את התרגיל.")
add_step(doc, 2, "הוסף גישה באתר.", "שייך מנהלים ומשתתפים רלוונטיים לתרגיל.")
add_step(doc, 3, "עקוב בצפייה חיה.", "סמן משתתפים, אירועים ושובל לפי הצורך.")
add_step(doc, 4, "דווח אירועים.", "הוסף תיאור, זמן ומיקום מהאפליקציה או מהאתר.")
add_step(doc, 5, "סגור את התרגיל.", "המשתתפים יקבלו הודעה והמעקב ייפסק.")
heading(doc, "תרחיש ג תחקור לאחר סיום", 2)
add_step(doc, 1, "בחר תרגיל סגור באתר.", "סמן את המשתתפים להשוואה.")
add_step(doc, 2, "הפעל את ציר הזמן.", "בחר קצב ושובל, והפעל או עצור לפי הצורך.")
add_step(doc, 3, "הצג אירועים ושמות.", "בדוק את הקשר בין האירועים לבין תנועת הכוחות.")
add_step(doc, 4, "ייצא וידאו.", "בחר טווח, איכות והגדרות והמתן להורדה המקומית.")

add_page_break(doc)
heading(doc, "5 פתרון תקלות", 1)
add_table(doc,
          ["בעיה", "מה לבדוק", "פעולה מומלצת"],
          [
              ["לא ניתן להיכנס", "מייל, סיסמה וחיבור לאינטרנט.", "נסה שכחתי סיסמה. אם אין הודעת הצלחה, פנה ל-ADMIN."],
              ["אין תרגילים פעילים", "התרגיל לא נסגר ושויכה לחשבון גישה.", "רענן את הרשימה; בקש ממנהל להוסיף גישה או לפתוח תרגיל."],
              ["אין נקודות GPS", "הרשאת מיקום, GPS פעיל ומצב מקליט GPS.", "אפשר מיקום מדויק וברקע ובטל הגבלת סוללה."],
              ["ממתינות נקודות לסנכרון", "חיבור אינטרנט וכתובת השרת.", "השאר את האפליקציה פעילה עד שהמונה יורד."],
              ["המסלול לא מופיע באתר", "המשתתף מסומן ויש לו נקודות.", "רענן, בחר מחדש את התרגיל וסמן את המשתתף."],
              ["אירוע לא מופיע", "הצג אירועים מסומן וזמן האירוע כבר הגיע.", "הפעל את הסימון והזז את ציר הזמן לאחר זמן האירוע."],
              ["צפייה חיה מושבתת", "סטטוס התרגיל.", "צפייה חיה זמינה רק בתרגיל ACTIVE."],
              ["וידאו לא נוצר", "Chrome או Edge, לשונית פעילה והרשאת הורדה.", "השאר את החלון פתוח, קצר את הטווח או בחר 720p."],
              ["האתר מציג 503", "Render מתעורר או נפרס מחדש.", "המתן כדקה ורענן. יצירת הווידאו המקומית אינה אמורה להעמיס על השרת."],
          ], widths=[1.55, 2.45, 2.7])
heading(doc, "מתי לפנות למנהל המערכת", 2)
add_bullet(doc, "כאשר סוג המשתמש או התפקידים הזמינים אינם נכונים.")
add_bullet(doc, "כאשר תרגיל קיים אינו מופיע למרות שהוא פעיל.")
add_bullet(doc, "כאשר חשבון ננעל, נמחק או אינו מקבל קישור לאיפוס.")
add_bullet(doc, "כאשר נתונים היסטוריים צריכים להיות משויכים לחשבון חדש.")

add_page_break(doc)
heading(doc, "6 כללי עבודה מומלצים", 1)
add_bullet(doc, "פתח תרגיל אחד בלבד לכל אירוע מבצעי והשתמש בשם ייחודי הכולל תאריך או אזור.")
add_bullet(doc, "בצע בדיקת GPS וסנכרון קצרה לפני תחילת התרגיל.")
add_bullet(doc, "אל תחליף חשבון באמצע תרגיל לפני שכל הנקודות הסתנכרנו.")
add_bullet(doc, "סגור תרגיל רק לאחר שכל המשתתפים סיימו והנתונים האחרונים הגיעו.")
add_bullet(doc, "מחק תרגיל רק לאחר תחקור, גיבוי או ייצוא נדרש.")
add_bullet(doc, "השתמש ב-ADMIN רק למשימות ניהול. לפעילות שוטפת עדיף MANAGER או USER.")
heading(doc, "מילון מונחים", 2)
add_table(doc,
          ["מונח", "הסבר"],
          [
              ["תרגיל פעיל", "תרגיל שניתן להצטרף אליו ולשלוח אליו נקודות ואירועים."],
              ["משתתף", "זהות התרגיל שמופיעה ברשימה ובמפה, עם שם ותפקיד."],
              ["נקודת GPS", "מיקום שנמדד בטלפון ונשמר עם חותמת זמן."],
              ["מסלול", "רצף נקודות GPS של משתתף אחד."],
              ["שובל", "מספר מוגבל של הנקודות האחרונות המוצגות בזמן הניגון."],
              ["אירוע", "דיווח הכולל תיאור, זמן, מיקום ושם ותפקיד המדווח."],
              ["צפייה חיה", "הצגת המיקום האחרון שנשמר לכל משתתף בתרגיל פעיל."],
              ["סנכרון", "העברת נתונים שנשמרו בטלפון אל השרת."],
              ["גישה לתרגיל", "הרשאה לחשבון מסוים לצפות או לנהל תרגיל מסוים."],
          ], widths=[1.7, 5.0])

add_page_break(doc)
heading(doc, "בדיקת מוכנות לפני תרגיל", 1)
for text in [
    "כל המשתמשים הצליחו להיכנס לחשבון הנכון.",
    "סוג המשתמש והתפקיד המבצעי נכונים.",
    "התרגיל מופיע כפעיל ברשימה.",
    "הרשאות מיקום, מיקום ברקע והתראות פעילות.",
    "במסך המעקב מופיע מקליט GPS.",
    "דיוק ה-GPS סביר ומונה הנקודות עולה.",
    "מונה ממתינות לסנכרון חוזר לאפס כאשר יש אינטרנט.",
    "המנהל רואה את המשתתפים ואת המסלולים באתר.",
    "דיווח אירוע ניסיוני מופיע במפה ובציר הזמן, אם התפקיד מאפשר זאת.",
]:
    add_text(doc, f"☐  {text}", space_after=5)
add_text(doc, "לאחר השלמת הבדיקה ניתן להתחיל את התרגיל. אם סעיף מהותי אינו תקין, מומלץ לעצור ולפתור אותו לפני יציאה לשטח.", bold_lead="לאחר השלמת הבדיקה")

# Final pass: normalize every paragraph and table cell to RTL, including built-in placeholders.
for paragraph in doc.paragraphs:
    rtl_paragraph(paragraph, paragraph.alignment or WD_ALIGN_PARAGRAPH.RIGHT)
    for run in paragraph.runs:
        rtl_run(run)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                rtl_paragraph(paragraph, paragraph.alignment or WD_ALIGN_PARAGRAPH.RIGHT)
                for run in paragraph.runs:
                    rtl_run(run)

doc.core_properties.title = "מדריך משתמש ReTrace"
doc.core_properties.subject = "מדריך מלא לאפליקציה ולמערכת הצפייה בתרגילים"
doc.core_properties.author = "ReTrace"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
